"""
Core analysis logic: parse Label Studio JSON export, generate summary report,
Excel workbook, YOLO training labels, and VLM fine-tuning JSONL.
"""

from __future__ import annotations

import glob as globmod
import json
import logging
import os
import re
import statistics
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Set, Tuple

from .config import AnalysisConfig
from .utils.ls_convert import interpolate_keyframes, ls_rect_to_yolo
from .utils.s3 import (
    meta_path_to_camera_date,
    meta_path_to_clip_id,
    meta_path_to_s3_clip,
    meta_path_to_s3_vlm_crop,
)

log = logging.getLogger("analysis")

DELETE_MARKER = "[delete]"


# ---------------------------------------------------------------------------
# Parsed data structures
# ---------------------------------------------------------------------------

@dataclass
class Track:
    """One bounding-box track (videorectangle) from an annotation."""
    task_id: int
    camera_name: str
    label: str
    num_keyframes: int
    first_frame: int
    last_frame: int
    frames_count: int
    duration: float
    avg_box_area_pct: float
    has_enabled_false: bool
    sequence: List[Dict[str, Any]]


@dataclass
class ParsedTask:
    """Structured data extracted from one Label Studio task + its annotations."""
    task_id: int
    camera_name: str
    kind: str
    date: str
    clip_start_local: str
    clip_end_local: str
    duration_sec: float
    meta_path: str
    has_vlm_crop: bool
    fps: float

    s3_clip_url: str
    s3_vlm_crop_url: str
    clip_id: str

    tracks: List[Track] = field(default_factory=list)
    vlm_description: str = ""
    annotator_id: Optional[int] = None
    lead_time_sec: float = 0.0
    num_annotations: int = 0

    # quality flags
    no_person: bool = False
    short_description: bool = False
    multi_annotator: bool = False
    low_keyframes: bool = False


# ---------------------------------------------------------------------------
# Parse export
# ---------------------------------------------------------------------------

def parse_export(
    export_path: str, cfg: AnalysisConfig,
) -> List[ParsedTask]:
    """Load and parse a Label Studio JSON export into structured task list."""
    log.info("Loading export from %s", export_path)
    with open(export_path, "r", encoding="utf-8") as f:
        raw_tasks: List[Dict[str, Any]] = json.load(f)
    log.info("Loaded %d tasks", len(raw_tasks))

    parsed: List[ParsedTask] = []
    for raw in raw_tasks:
        data = raw.get("data", {})
        meta_path = data.get("meta_path", "")
        camera, date = meta_path_to_camera_date(meta_path)

        pt = ParsedTask(
            task_id=raw.get("id", 0),
            camera_name=data.get("camera_name", camera),
            kind=data.get("kind", "unknown"),
            date=date,
            clip_start_local=data.get("clip_start_local", ""),
            clip_end_local=data.get("clip_end_local", ""),
            duration_sec=float(data.get("duration_sec", 0)),
            meta_path=meta_path,
            has_vlm_crop=bool(data.get("has_vlm_crop", False)),
            fps=float(data.get("fps", 7)),
            s3_clip_url=meta_path_to_s3_clip(meta_path, cfg.s3_bucket, cfg.s3_prefix),
            s3_vlm_crop_url=meta_path_to_s3_vlm_crop(meta_path, cfg.s3_bucket, cfg.s3_prefix),
            clip_id=meta_path_to_clip_id(meta_path),
        )

        annotations = raw.get("annotations", [])
        pt.num_annotations = len(annotations)
        pt.multi_annotator = len(annotations) > 1

        annotator_ids = set()
        lead_times = []
        for ann in annotations:
            if ann.get("was_cancelled"):
                continue
            cby = ann.get("completed_by")
            if cby is not None:
                annotator_ids.add(cby)
            lt = ann.get("lead_time")
            if lt is not None:
                lead_times.append(float(lt))

            for r in ann.get("result", []):
                rtype = r.get("type")
                val = r.get("value", {})

                if rtype == "videorectangle":
                    labels = val.get("labels", [])
                    seq = val.get("sequence", [])
                    frames_count = int(val.get("framesCount", 0))
                    duration = float(val.get("duration", 0))

                    enabled_kfs = [kf for kf in seq if kf.get("enabled", True)]
                    frame_nums = [int(kf.get("frame", 0)) for kf in enabled_kfs]

                    areas = []
                    for kf in enabled_kfs:
                        areas.append(kf.get("width", 0) * kf.get("height", 0))

                    track = Track(
                        task_id=pt.task_id,
                        camera_name=pt.camera_name,
                        label=labels[0] if labels else "unknown",
                        num_keyframes=len(seq),
                        first_frame=min(frame_nums) if frame_nums else 0,
                        last_frame=max(frame_nums) if frame_nums else 0,
                        frames_count=frames_count,
                        duration=duration,
                        avg_box_area_pct=statistics.mean(areas) if areas else 0,
                        has_enabled_false=any(not kf.get("enabled", True) for kf in seq),
                        sequence=seq,
                    )
                    pt.tracks.append(track)

                elif rtype == "textarea" and r.get("from_name") == "vlm_description":
                    texts = val.get("text", [])
                    if texts:
                        pt.vlm_description = texts[0]

        if annotator_ids:
            pt.annotator_id = min(annotator_ids)
        if lead_times:
            pt.lead_time_sec = sum(lead_times) / len(lead_times)

        # quality flags
        person_count = sum(1 for t in pt.tracks if t.label == "person")
        pt.no_person = person_count == 0
        pt.short_description = len(pt.vlm_description.strip()) < 10
        min_expected_kfs = max(3, int(pt.duration_sec * pt.fps * 0.1))
        pt.low_keyframes = any(
            t.num_keyframes < min_expected_kfs for t in pt.tracks
        )

        parsed.append(pt)

    log.info("Parsed %d tasks with %d total tracks",
             len(parsed), sum(len(p.tracks) for p in parsed))
    return parsed


# ---------------------------------------------------------------------------
# [delete] marker filtering
# ---------------------------------------------------------------------------

def _is_delete_marker(description: str) -> bool:
    return description.strip().lower() == DELETE_MARKER


def _derive_s3_keys(meta_path: str, prefix: str) -> List[str]:
    """Derive all known S3 keys for a clip from its meta_path."""
    rel = meta_path.replace("\\", "/")
    camera_date = "/".join(rel.split("/")[1:3])  # e.g. "main_door/2026-03-02"
    clip_id = re.sub(r"\.meta\.json$", "", os.path.basename(rel))

    keys = [
        f"{prefix}/clips/{camera_date}/{clip_id}.mp4",
        f"{prefix}/meta/{camera_date}/{clip_id}.meta.json",
        f"{prefix}/vlm_crops/{camera_date}/{clip_id}.mp4",
        f"{prefix}/responses/{camera_date}/{clip_id}.model_raw.txt",
    ]
    return keys


def _delete_s3_objects(
    bucket: str, prefix: str, tasks: List[ParsedTask],
) -> int:
    """Delete all S3 objects associated with *tasks*. Returns count deleted."""
    try:
        import boto3
    except ImportError:
        log.warning("boto3 not installed — skipping S3 deletion.")
        return 0

    s3 = boto3.client("s3")
    all_keys: List[str] = []

    for pt in tasks:
        all_keys.extend(_derive_s3_keys(pt.meta_path, prefix))

        rel = pt.meta_path.replace("\\", "/")
        camera_date = "/".join(rel.split("/")[1:3])
        paginator = s3.get_paginator("list_objects_v2")
        for subdir in ("yolo/images", "yolo/labels"):
            yolo_prefix = f"{prefix}/{subdir}/{camera_date}/{pt.clip_id}_f"
            for page in paginator.paginate(Bucket=bucket, Prefix=yolo_prefix):
                for obj in page.get("Contents", []):
                    all_keys.append(obj["Key"])

    all_keys = list(set(all_keys))
    if not all_keys:
        return 0

    log.info("Deleting %d S3 objects for %d marked tasks...", len(all_keys), len(tasks))
    deleted = 0
    batch_size = 1000
    for i in range(0, len(all_keys), batch_size):
        batch = all_keys[i : i + batch_size]
        objects = [{"Key": k} for k in batch]
        resp = s3.delete_objects(Bucket=bucket, Delete={"Objects": objects, "Quiet": True})
        deleted += len(batch) - len(resp.get("Errors", []))
        for err in resp.get("Errors", []):
            log.warning("S3 delete error: %s — %s", err["Key"], err["Message"])

    return deleted


def _delete_local_files(dataset_dir: str, tasks: List[ParsedTask]) -> int:
    """Delete local files associated with *tasks*. Returns count deleted."""
    if not dataset_dir or not os.path.isdir(dataset_dir):
        return 0

    deleted = 0
    for pt in tasks:
        rel = pt.meta_path.replace("\\", "/")
        camera_date = "/".join(rel.split("/")[1:3])
        clip_id = pt.clip_id

        candidates = [
            os.path.join(dataset_dir, "clips", camera_date, f"{clip_id}.mp4"),
            os.path.join(dataset_dir, "meta", camera_date, f"{clip_id}.meta.json"),
            os.path.join(dataset_dir, "vlm_crops", camera_date, f"{clip_id}.mp4"),
            os.path.join(dataset_dir, "responses", camera_date, f"{clip_id}.model_raw.txt"),
        ]

        for subdir in ("yolo/images", "yolo/labels"):
            pattern = os.path.join(dataset_dir, subdir, camera_date, f"{clip_id}_f*")
            candidates.extend(globmod.glob(pattern))

        for fp in candidates:
            if os.path.isfile(fp):
                try:
                    os.remove(fp)
                    deleted += 1
                except OSError as exc:
                    log.warning("Failed to delete %s: %s", fp, exc)

    return deleted


def filter_deleted_tasks(
    tasks: List[ParsedTask],
    cfg: AnalysisConfig,
    dataset_dir: Optional[str] = None,
) -> Tuple[List[ParsedTask], List[ParsedTask]]:
    """
    Partition *tasks* into (keep, deleted) based on the ``[delete]`` VLM marker.

    For deleted tasks: removes associated files from S3 and local disk.
    """
    keep: List[ParsedTask] = []
    deleted: List[ParsedTask] = []

    for pt in tasks:
        if _is_delete_marker(pt.vlm_description):
            deleted.append(pt)
        else:
            keep.append(pt)

    if not deleted:
        log.info("No tasks marked %s.", DELETE_MARKER)
        return keep, deleted

    log.info(
        "Found %d tasks marked %s (keeping %d).",
        len(deleted), DELETE_MARKER, len(keep),
    )

    s3_deleted = _delete_s3_objects(cfg.s3_bucket, cfg.s3_prefix, deleted)
    log.info("Deleted %d S3 objects for %d marked tasks.", s3_deleted, len(deleted))

    local_deleted = _delete_local_files(dataset_dir, deleted) if dataset_dir else 0
    if local_deleted:
        log.info("Deleted %d local files for %d marked tasks.", local_deleted, len(deleted))

    return keep, deleted


# ---------------------------------------------------------------------------
# Summary report
# ---------------------------------------------------------------------------

def build_summary_report(
    tasks: List[ParsedTask],
    cfg: AnalysisConfig,
    deleted_tasks: Optional[List[ParsedTask]] = None,
) -> str:
    """Generate a markdown summary report string."""
    lines: List[str] = []

    def h1(s: str) -> None:
        lines.append(f"# {s}\n")

    def h2(s: str) -> None:
        lines.append(f"## {s}\n")

    def h3(s: str) -> None:
        lines.append(f"### {s}\n")

    def row(k: str, v: Any) -> None:
        lines.append(f"- **{k}**: {v}")

    all_tracks = [t for p in tasks for t in p.tracks]
    all_vlm = [p.vlm_description for p in tasks if p.vlm_description]

    h1("Label Export Analysis Report")
    lines.append("")

    # --- Overview ---
    h2("Overview")
    row("Total tasks", len(tasks))
    row("Total annotations", sum(p.num_annotations for p in tasks))
    row("Total bounding-box tracks", len(all_tracks))
    row("Tasks with VLM descriptions", len(all_vlm))
    durations = [p.duration_sec for p in tasks]
    if durations:
        row("Duration range",
            f"{min(durations):.1f}s - {max(durations):.1f}s "
            f"(mean {statistics.mean(durations):.1f}s)")
    row("Total video time", f"{sum(durations):.0f}s ({sum(durations)/60:.1f} min)")
    lines.append("")

    # --- Per-camera ---
    h2("Per-Camera Breakdown")
    cam_tasks: Dict[str, List[ParsedTask]] = defaultdict(list)
    for p in tasks:
        cam_tasks[p.camera_name].append(p)
    for cam in sorted(cam_tasks):
        pts = cam_tasks[cam]
        cam_tracks = [t for p in pts for t in p.tracks]
        persons = sum(1 for t in cam_tracks if t.label == "person")
        cars = sum(1 for t in cam_tracks if t.label == "car")
        other = len(cam_tracks) - persons - cars
        h3(cam)
        row("Tasks", len(pts))
        row("Tracks", f"{len(cam_tracks)} (person: {persons}, car: {cars}, other: {other})")
        row("Kinds", dict(Counter(p.kind for p in pts)))
        lines.append("")

    # --- Label distribution ---
    h2("Label Distribution")
    label_counts = Counter(t.label for t in all_tracks)
    for label, count in label_counts.most_common():
        row(label, count)
    lines.append("")

    # --- Annotator stats ---
    h2("Annotator Statistics")
    ann_tasks: Dict[int, List[ParsedTask]] = defaultdict(list)
    for p in tasks:
        if p.annotator_id is not None:
            ann_tasks[p.annotator_id].append(p)
    for aid in sorted(ann_tasks):
        pts = ann_tasks[aid]
        lead_times = [p.lead_time_sec for p in pts if p.lead_time_sec > 0]
        h3(f"Annotator {aid}")
        row("Tasks annotated", len(pts))
        if lead_times:
            row("Avg lead time", f"{statistics.mean(lead_times):.0f}s")
            row("Total time spent", f"{sum(lead_times):.0f}s ({sum(lead_times)/60:.1f} min)")
        lines.append("")

    # --- VLM descriptions ---
    h2("VLM Description Statistics")
    if all_vlm:
        lengths = [len(v) for v in all_vlm]
        row("Count", len(all_vlm))
        row("Avg length", f"{statistics.mean(lengths):.0f} chars")
        row("Min / Max length", f"{min(lengths)} / {max(lengths)} chars")
        lines.append("")
        h3("Most Common Descriptions")
        desc_counts = Counter(all_vlm)
        for desc, count in desc_counts.most_common(10):
            truncated = desc[:120] + ("..." if len(desc) > 120 else "")
            lines.append(f"- [{count}x] {truncated}")
    else:
        row("Count", 0)
    lines.append("")

    # --- Keyframe density ---
    h2("Annotation Quality Metrics")
    kf_counts = [t.num_keyframes for t in all_tracks]
    if kf_counts:
        row("Keyframes per track",
            f"min={min(kf_counts)}, max={max(kf_counts)}, "
            f"mean={statistics.mean(kf_counts):.1f}, "
            f"median={statistics.median(kf_counts):.1f}")
    tracks_per_task = [len(p.tracks) for p in tasks]
    if tracks_per_task:
        row("Tracks per task",
            f"min={min(tracks_per_task)}, max={max(tracks_per_task)}, "
            f"mean={statistics.mean(tracks_per_task):.1f}")
    enabled_false = sum(1 for t in all_tracks if t.has_enabled_false)
    row("Tracks with visibility gaps (enabled=False)", enabled_false)
    lines.append("")

    # --- Quality flags ---
    h2("Quality Flags (tasks to review)")
    no_person = [p for p in tasks if p.no_person]
    short_desc = [p for p in tasks if p.short_description]
    multi_ann = [p for p in tasks if p.multi_annotator]
    low_kf = [p for p in tasks if p.low_keyframes]
    row("No person detected (trigger tasks)", len(no_person))
    row("Short VLM description (<10 chars)", len(short_desc))
    row("Multiple annotators (needs review)", len(multi_ann))
    row("Low keyframe density", len(low_kf))
    if no_person:
        lines.append(f"  - No-person task IDs: {[p.task_id for p in no_person]}")
    if short_desc:
        lines.append(f"  - Short-description task IDs: {[p.task_id for p in short_desc]}")
    if multi_ann:
        lines.append(f"  - Multi-annotator task IDs: {[p.task_id for p in multi_ann]}")
    lines.append("")

    # --- Deleted tasks ---
    if deleted_tasks:
        h2(f"Deleted Tasks ({DELETE_MARKER} marker)")
        row("Total deleted", len(deleted_tasks))
        del_cams: Dict[str, int] = Counter(p.camera_name for p in deleted_tasks)
        row("By camera", dict(del_cams))
        lines.append("")
        lines.append("| Task ID | Camera | Clip ID |")
        lines.append("|---------|--------|---------|")
        for p in deleted_tasks:
            lines.append(f"| {p.task_id} | {p.camera_name} | {p.clip_id} |")
        lines.append("")

    return "\n".join(lines)


def write_summary_report(
    tasks: List[ParsedTask],
    cfg: AnalysisConfig,
    output_dir: str,
    deleted_tasks: Optional[List[ParsedTask]] = None,
) -> str:
    """Build summary, print to console, save to file. Returns the report text."""
    report = build_summary_report(tasks, cfg, deleted_tasks=deleted_tasks)
    print(report)

    os.makedirs(output_dir, exist_ok=True)
    report_path = os.path.join(output_dir, "summary_report.md")
    with open(report_path, "w", encoding="utf-8") as f:
        f.write(report)
    log.info("Summary report saved to %s", report_path)
    return report


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def write_excel(
    tasks: List[ParsedTask], cfg: AnalysisConfig, output_dir: str,
) -> str:
    """Write an Excel workbook with task-level and track-level sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    os.makedirs(output_dir, exist_ok=True)
    wb = Workbook()

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    wrap_align = Alignment(wrap_text=True, vertical="top")

    # ── Sheet 1: Tasks ─────────────────────────────────────────────────
    ws_tasks = wb.active
    ws_tasks.title = "Tasks"

    task_headers = [
        "task_id", "camera_name", "kind", "date",
        "clip_start_local", "clip_end_local", "duration_sec",
        "meta_path", "s3_clip_url", "s3_vlm_crop_url",
        "has_vlm_crop",
        "num_person_tracks", "num_car_tracks", "num_other_tracks",
        "total_tracks", "total_keyframes",
        "vlm_description",
        "annotator_id", "lead_time_sec", "num_annotations",
        "flag_no_person", "flag_short_description",
        "flag_multi_annotator", "flag_low_keyframes",
    ]
    for col, header in enumerate(task_headers, 1):
        cell = ws_tasks.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, pt in enumerate(tasks, 2):
        persons = sum(1 for t in pt.tracks if t.label == "person")
        cars = sum(1 for t in pt.tracks if t.label == "car")
        other = len(pt.tracks) - persons - cars
        total_kfs = sum(t.num_keyframes for t in pt.tracks)

        values = [
            pt.task_id, pt.camera_name, pt.kind, pt.date,
            pt.clip_start_local, pt.clip_end_local, round(pt.duration_sec, 2),
            pt.meta_path, pt.s3_clip_url, pt.s3_vlm_crop_url,
            pt.has_vlm_crop,
            persons, cars, other,
            len(pt.tracks), total_kfs,
            pt.vlm_description,
            pt.annotator_id, round(pt.lead_time_sec, 1), pt.num_annotations,
            pt.no_person, pt.short_description,
            pt.multi_annotator, pt.low_keyframes,
        ]
        for col, val in enumerate(values, 1):
            cell = ws_tasks.cell(row=row_idx, column=col, value=val)
            if col == task_headers.index("vlm_description") + 1:
                cell.alignment = wrap_align

    for col in range(1, len(task_headers) + 1):
        ws_tasks.column_dimensions[
            ws_tasks.cell(row=1, column=col).column_letter
        ].width = 18
    ws_tasks.column_dimensions["Q"].width = 50  # vlm_description
    ws_tasks.column_dimensions["I"].width = 60  # s3_clip_url
    ws_tasks.column_dimensions["J"].width = 60  # s3_vlm_crop_url

    ws_tasks.auto_filter.ref = ws_tasks.dimensions

    # ── Sheet 2: Tracks ────────────────────────────────────────────────
    ws_tracks = wb.create_sheet("Tracks")

    track_headers = [
        "task_id", "camera_name", "track_label",
        "num_keyframes", "first_frame", "last_frame", "frames_count",
        "avg_box_area_pct", "has_visibility_gaps",
    ]
    for col, header in enumerate(track_headers, 1):
        cell = ws_tracks.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    row_idx = 2
    for pt in tasks:
        for trk in pt.tracks:
            values = [
                trk.task_id, trk.camera_name, trk.label,
                trk.num_keyframes, trk.first_frame, trk.last_frame,
                trk.frames_count,
                round(trk.avg_box_area_pct, 2), trk.has_enabled_false,
            ]
            for col, val in enumerate(values, 1):
                ws_tracks.cell(row=row_idx, column=col, value=val)
            row_idx += 1

    for col in range(1, len(track_headers) + 1):
        ws_tracks.column_dimensions[
            ws_tracks.cell(row=1, column=col).column_letter
        ].width = 18
    ws_tracks.auto_filter.ref = ws_tracks.dimensions

    excel_path = os.path.join(output_dir, "analysis.xlsx")
    wb.save(excel_path)
    log.info("Excel workbook saved to %s (%d task rows, %d track rows)",
             excel_path, len(tasks), row_idx - 2)
    return excel_path


# ---------------------------------------------------------------------------
# YOLO training data export
# ---------------------------------------------------------------------------

def _label_to_class_id(label: str, coco_labels: Dict[int, str]) -> Optional[int]:
    """Reverse lookup: label name -> COCO class ID."""
    for cid, name in coco_labels.items():
        if name == label:
            return cid
    return None


def write_yolo_labels(
    tasks: List[ParsedTask], cfg: AnalysisConfig, output_dir: str,
) -> str:
    """
    Convert LS videorectangle annotations to per-frame YOLO label files.

    Interpolates between keyframes to produce a bounding box for every frame,
    then writes one .txt label file per frame.  Class IDs are remapped to
    sequential 0..N-1 matching the generated data.yaml.
    """
    yolo_dir = os.path.join(output_dir, "yolo")
    labels_dir = os.path.join(yolo_dir, "labels")
    os.makedirs(labels_dir, exist_ok=True)

    # Build class-ID remap: sparse COCO id -> sequential 0..N-1
    sorted_coco_ids = sorted(cfg.coco_labels.keys())
    coco_to_seq: Dict[int, int] = {
        cid: idx for idx, cid in enumerate(sorted_coco_ids)
    }
    names_dict = {idx: cfg.coco_labels[cid] for idx, cid in enumerate(sorted_coco_ids)}

    total_files = 0
    total_boxes = 0
    skipped_labels: Counter = Counter()

    for pt in tasks:
        cam_date_dir = os.path.join(labels_dir, pt.camera_name, pt.date)
        os.makedirs(cam_date_dir, exist_ok=True)

        frame_boxes: Dict[int, List[str]] = defaultdict(list)

        for trk in pt.tracks:
            coco_id = _label_to_class_id(trk.label, cfg.coco_labels)
            if coco_id is None:
                skipped_labels[trk.label] += 1
                continue
            seq_id = coco_to_seq[coco_id]

            frames_count = trk.frames_count or int(pt.duration_sec * pt.fps)
            per_frame = interpolate_keyframes(trk.sequence, frames_count)

            for frame_num, box in per_frame.items():
                xc, yc, w, h = ls_rect_to_yolo(
                    box["x"], box["y"], box["width"], box["height"],
                )
                if w <= 0 or h <= 0:
                    continue
                line = f"{seq_id} {xc:.6f} {yc:.6f} {w:.6f} {h:.6f}"
                frame_boxes[frame_num].append(line)
                total_boxes += 1

        for frame_num in sorted(frame_boxes):
            fname = f"{pt.clip_id}_f{frame_num:04d}.txt"
            fpath = os.path.join(cam_date_dir, fname)
            with open(fpath, "w", encoding="utf-8") as f:
                f.write("\n".join(frame_boxes[frame_num]) + "\n")
            total_files += 1

    if skipped_labels:
        log.warning("Skipped labels not in COCO map: %s", dict(skipped_labels))

    # Write data.yaml for Ultralytics
    import yaml
    data_yaml = {
        "path": os.path.abspath(yolo_dir),
        "train": "labels",
        "val": "labels",
        "names": names_dict,
        "nc": len(names_dict),
    }
    data_yaml_path = os.path.join(yolo_dir, "data.yaml")
    with open(data_yaml_path, "w", encoding="utf-8") as f:
        yaml.safe_dump(data_yaml, f, default_flow_style=False, sort_keys=False)

    log.info("YOLO export: %d label files, %d total boxes -> %s",
             total_files, total_boxes, labels_dir)
    return yolo_dir


# ---------------------------------------------------------------------------
# VLM fine-tuning JSONL export
# ---------------------------------------------------------------------------

def write_vlm_jsonl(
    tasks: List[ParsedTask], cfg: AnalysisConfig, output_dir: str,
) -> str:
    """Write a JSONL file for VLM fine-tuning (one line per task)."""
    os.makedirs(output_dir, exist_ok=True)
    jsonl_path = os.path.join(output_dir, "vlm_training.jsonl")

    count = 0
    with open(jsonl_path, "w", encoding="utf-8") as f:
        for pt in tasks:
            if not pt.vlm_description.strip():
                continue

            persons = sum(1 for t in pt.tracks if t.label == "person")
            cars = sum(1 for t in pt.tracks if t.label == "car")

            record = {
                "video_s3_path": pt.s3_clip_url,
                "vlm_crop_s3_path": pt.s3_vlm_crop_url,
                "description": pt.vlm_description,
                "camera_name": pt.camera_name,
                "duration_sec": round(pt.duration_sec, 2),
                "num_persons": persons,
                "num_cars": cars,
                "kind": pt.kind,
                "date": pt.date,
                "clip_id": pt.clip_id,
            }
            f.write(json.dumps(record, ensure_ascii=False) + "\n")
            count += 1

    log.info("VLM JSONL: %d records -> %s", count, jsonl_path)
    return jsonl_path


# ---------------------------------------------------------------------------
# Top-level runner
# ---------------------------------------------------------------------------

def run(
    export_path: str,
    cfg: AnalysisConfig,
    output_dir: str,
    *,
    dataset_dir: Optional[str] = None,
    skip_report: bool = False,
    skip_excel: bool = False,
    skip_yolo: bool = False,
    skip_vlm: bool = False,
) -> None:
    """Run the full analysis pipeline."""
    tasks = parse_export(export_path, cfg)

    if not tasks:
        log.warning("No tasks found in export file.")
        return

    tasks, deleted_tasks = filter_deleted_tasks(tasks, cfg, dataset_dir)

    if not tasks:
        log.warning("All tasks were marked %s. Nothing to analyze.", DELETE_MARKER)
        return

    if not skip_report:
        write_summary_report(tasks, cfg, output_dir, deleted_tasks=deleted_tasks)

    if not skip_excel:
        write_excel(tasks, cfg, output_dir)

    if not skip_yolo:
        write_yolo_labels(tasks, cfg, output_dir)

    if not skip_vlm:
        write_vlm_jsonl(tasks, cfg, output_dir)

    log.info("Analysis complete. Output directory: %s", output_dir)
