"""Publish a Studio collection as a tagging batch: s3 `tagging/<batch>/` in exactly the layout of the founder's
Label Studio batches (tagging/ameer_house_batch_1, _2, ...), so `home_guard_project/analysis` reads it unchanged.

Layout (mirror of tagging/ameer_house_batch_2):

    tagging/<batch>/<batch>.json                  Label Studio export JSON: tasks with their annotation and
                                                  predictions (analysis/analyze.py parse_export reads it)
    tagging/<batch>/dataset_multi/{clips,meta,responses,vlm_crops,yolo/images,yolo/labels}/<camera>/<date>/...
                                                  server-side copies of each clip's box files (same relative paths)
    tagging/<batch>/dataset_multi/label_studio_config.xml   the XML labeling/tasks.py builds
    tagging/<batch>/dataset_multi/label_studio_tasks.json   the tasks without annotations
    tagging/<batch>/analysis_output/yolo/{images,labels}/<camera>/<date>/<clip_id>_f<NNNN>.{jpg,txt}
                                                  EVERY frame (NNNN = native 0-based index), labels from box_at
                                                  with contiguous class ids, an empty file = a human negative
    tagging/<batch>/analysis_output/yolo/data.yaml
    tagging/<batch>/analysis_output/vlm_training.jsonl, summary_report.md, analysis.xlsx
    tagging/<batch>/_admin_center.json            marker: this batch was written by the Admin Center

Only clips whose current annotation is submitted or reviewed are published (the others are listed as missing,
"not labeled"); a clip labeled `drop_clip` stays in the JSON (its text is the `[delete]` marker analyze.py honours)
but is left out of analysis_output. The existing Label Studio batches are read-only forever: a name that exists on
S3 without the marker is refused, and the build writes only under `tagging/<batch>/` (a scoped S3 writer).

One attempt per batch name, forever: `tagging_publishes.batch_name` is unique across all states and a folder that
already exists is never written again, so there is no resume. Marker states: building -> ready | partial | failed;
a failed attempt keeps its folder (marker "failed", best effort) and the name is spent. Because a name is never
reused, a worker that lost its lease can only ever write into its own attempt's folder. The worker checks it still
owns the job (worker id + running state) before and after each clip and every FENCE_EVERY frames; it becomes
ready/partial with one conditional UPDATE (this worker, running, fresh heartbeat) BEFORE the final marker is
written, and writes nothing more when that UPDATE does not match.

Frames are in the clip's NATIVE frame space (`data.frame_space` = "native"): `data.fps` is the clip's real decoded
fps, a keyframe on native frame n is Label Studio frame n + 1 (LS frames are 1-based), `framesCount` and
`data.frames_count` are the decoded native frame count and `time` the keyframe's real timestamp. Only constant-rate
clips are published: every decoded frame's timestamp must be within half a frame of index / fps (the time the
analysis pipeline assumes), else the clip is missing "variable frame rate"; a clip without a known fps is missing
"frame rate unknown". Image and label files use the native 0-based index; labels are box_at(index / fps).
"""
from __future__ import annotations

import io
import json
import logging
import math
import re
import statistics
import tempfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Optional

from sqlalchemy import func, or_, select, text, update
from sqlalchemy.orm import Session

from home_guard_project.fleet_contract.classes import COCO_NAMES

from . import labeling, media
from .models import (Annotation, Artifact, Collection, CollectionItem, Customer, Device, Event, Staff,
                     TaggingPublish)
from .s3 import ETagMismatch

log = logging.getLogger(__name__)

ROOT = "tagging/"
MARKER = "_admin_center.json"
PROTECTED = frozenset({"ameer_house_batch_1", "ameer_house_batch_2", "uca_dataset_batch", "smarthome_dataset_batch"})
NAME = re.compile(r"[a-z0-9_]{1,64}")
FRAME_SPACE = "native"
DELETE_MARKER = "[delete]"
NOT_LABELED = "not labeled"
NO_CONSENT = "no training consent"
SOURCE_MISSING = "source missing"
VARIABLE_RATE = "variable frame rate"
RATE_UNKNOWN = "frame rate unknown"
PUBLISHED = "This batch is published; choose a new name"
FAILED_NAME = "This name was used by a failed publish; choose a new name"
# budget (estimated before anything is queued)
MAX_FRAMES = 300_000
MAX_BYTES = 30 * 10 ** 9
JPEG_BYTES_PER_PIXEL = 0.25  # -q:v 2 JPEGs: ~0.5 MB for a 1080p frame
FALLBACK_JPEG_BYTES = 200_000
FALLBACK_FRAMES = 300  # a clip whose length is unknown
LABEL_BYTES = 64
FENCE_EVERY = 100  # frames between ownership checks inside one clip
_COPY_ROLES = ("original_video", "meta", "raw_answer", "teacher_frame", "crop_video", "yolo_image", "yolo_label")
_MIME = {".mp4": "video/mp4", ".json": "application/json", ".txt": "text/plain", ".jpg": "image/jpeg",
         ".jpeg": "image/jpeg", ".png": "image/png"}
LABEL_NAMES = sorted(COCO_NAMES.values())  # labeling/config get_all_label_names(): alphabetical

# labeling/tasks.py _LABEL_CONFIG_XML, verbatim
_LABEL_CONFIG_XML = """\
<View>
  <Header value="$header_info"/>
  <View style="display:flex; gap:16px;">
    <View style="flex:1;">
      <Header value="Full Frame — YOLO Re-tagging"/>
      <Labels name="label" toName="video_full">
{label_tags}
      </Labels>
      <Video name="video_full" value="$video_url" frameRate="$fps"/>
      <VideoRectangle name="bbox" toName="video_full"/>
    </View>
    <View style="flex:1;">
      <Header value="$vlm_crop_header"/>
      <Video name="video_crop" value="$vlm_crop_url" frameRate="$vlm_fps"/>
      <Header value="Scene Description (for VLM training)"/>
      <TextArea name="vlm_description" toName="video_crop"
                rows="4" editable="true"
                placeholder="Describe what happens in the video..."/>
    </View>
  </View>
</View>
"""


def label_config_xml() -> str:
    """The Label Studio labeling config labeling/tasks.py build_label_config() writes."""
    return _LABEL_CONFIG_XML.format(label_tags="\n".join(f'    <Label value="{n}"/>' for n in LABEL_NAMES))


def batch_prefix(name: str) -> str:
    return f"{ROOT}{name}/"


class PublishRefused(Exception):
    """A publish request the server refuses (HTTP 409); the message is shown as is."""


class PublishTooBig(Exception):
    """A publish over the frame/byte budget (HTTP 400); the message carries the estimate."""


# ---------------------------------------------------------------- request

def _check_format(name: str) -> None:
    if not NAME.fullmatch(name or ""):
        raise PublishRefused("A batch name is 1-64 characters: a-z, 0-9 and _")
    if name in PROTECTED:
        raise PublishRefused(f"tagging/{name}/ is a Label Studio batch and is read-only")


def _check_folder(s3, name: str) -> str:
    """The batch prefix, which must not exist yet; PublishRefused for any folder already there (a name is used
    once: a published or partial batch, a failed or interrupted attempt, or a folder not made here)."""
    prefix = batch_prefix(name)
    if not s3.any_under(prefix):
        return prefix
    if not s3.exists(prefix + MARKER):
        raise PublishRefused(f"tagging/{name}/ already exists and was not written by the Admin Center; "
                             "choose another name")
    try:
        marker = json.loads(s3.get_bytes(prefix + MARKER))
    except (ValueError, UnicodeDecodeError):
        marker = None
    if not isinstance(marker, dict):
        raise PublishRefused(f"tagging/{name}/ has an unreadable marker; choose another name")
    raise PublishRefused(PUBLISHED if marker.get("state") in ("ready", "partial") else FAILED_NAME)


def _check_unused(session: Session, name: str) -> None:
    """PublishRefused when a publish of this name exists in any state (the name is unique, forever)."""
    state = session.scalar(select(TaggingPublish.state).where(TaggingPublish.batch_name == name))
    if state is None:
        return
    if state in ("queued", "running"):
        raise PublishRefused(f"A publish of {name} is already running")
    raise PublishRefused(PUBLISHED if state in ("ready", "partial") else FAILED_NAME)


def frame_rate_problem(times: list[float], fps) -> Optional[str]:
    """None for a constant-rate clip: every decoded frame n has a timestamp within half a frame of n / fps (the
    time the analysis pipeline gives native frame n). Else RATE_UNKNOWN (no usable fps) or VARIABLE_RATE (a frame
    off its slot, or one without a timestamp)."""
    if isinstance(fps, bool) or not isinstance(fps, (int, float)) or not math.isfinite(fps) or fps <= 0:
        return RATE_UNKNOWN
    half = 0.5 / fps
    for n, t in enumerate(times):
        if not isinstance(t, (int, float)) or not math.isfinite(t) or abs(t - n / fps) > half:
            return VARIABLE_RATE
    return None


def take_snapshot(session: Session, collection_id: int) -> dict:
    """Which clips the batch will hold, frozen at request time: every event of the collection whose current
    annotation is submitted or reviewed (with that version); the others are `missing`."""
    from .models import AnnotationHead

    rows = session.execute(
        select(Event.id, Customer.consent_training, AnnotationHead.version, AnnotationHead.status,
               AnnotationHead.drop_clip)
        .join(CollectionItem, CollectionItem.event_id == Event.id)
        .join(Device, Device.id == Event.device_pk).join(Customer, Customer.id == Device.customer_id)
        .outerjoin(AnnotationHead, AnnotationHead.event_id == Event.id)
        .where(CollectionItem.collection_id == collection_id).order_by(Event.id)).all()
    events, missing = [], []
    for eid, consent, version, status, drop in rows:
        if not consent:
            missing.append({"event_id": eid, "reason": NO_CONSENT})
        elif status not in labeling.DONE:
            missing.append({"event_id": eid, "reason": NOT_LABELED})
        else:
            events.append({"id": eid, "version": version, "status": status, "drop_clip": bool(drop)})
    return {"events": events, "missing": missing}


def _sources(session: Session, ev_id: int) -> dict[str, Artifact]:
    """{relative path under the box root: artifact} of the clip's box files, training copy first."""
    out: dict[str, Artifact] = {}
    arts = session.scalars(select(Artifact).where(Artifact.event_id == ev_id, Artifact.role.in_(_COPY_ROLES),
                                                  Artifact.available.is_(True))).all()
    for art in sorted(arts, key=lambda a: (not a.s3_key.startswith("dataset_"), a.s3_key)):
        rel = art.s3_key.split("/", 1)[1] if "/" in art.s3_key else None
        if rel and rel not in out:
            out[rel] = art
    return out


def estimate(session: Session, snapshot: dict) -> tuple[int, int]:
    """(frames, bytes) the batch would write: every frame of each kept clip (a JPEG sized by the clip's frame size,
    plus its label) and the copies of the clips' box files."""
    frames = size = 0
    for info in snapshot.get("events", []):
        ev = session.get(Event, info["id"])
        if ev is None:
            continue
        size += sum(a.bytes or 0 for a in _sources(session, ev.id).values())
        if info.get("drop_clip"):
            continue
        fps, count, duration = labeling.clip_timing(ev, labeling.meta_body(session, ev.id))
        n = count or (int(round(duration * fps)) if duration and fps else FALLBACK_FRAMES)
        fs = ev.frame_size if isinstance(ev.frame_size, list) and len(ev.frame_size) == 2 else None
        per = (int(fs[0] * fs[1] * JPEG_BYTES_PER_PIXEL)
               if fs and all(isinstance(v, int) and v > 0 for v in fs) else FALLBACK_JPEG_BYTES)
        frames += n
        size += n * (per + LABEL_BYTES)
    return frames, size


def _check_budget(frames: int, size: int) -> None:
    if frames > MAX_FRAMES or size > MAX_BYTES:
        raise PublishTooBig(
            f"This batch is too big to publish: about {frames:,} frames and {size / 1e9:.1f} GB "
            f"(the limits are {MAX_FRAMES:,} frames and {MAX_BYTES / 1e9:.0f} GB). Split the collection.")


def create_publish(session: Session, s3, col: Collection, batch_name: str, staff: Staff,
                   now: datetime) -> TaggingPublish:
    """A queued publish of `col` as `batch_name`; PublishRefused (read-only batch, a name already used by any
    publish, or an existing folder) or PublishTooBig (over the frame/byte budget). The caller commits and hands
    the job to the worker pool."""
    _check_format(batch_name)
    session.execute(text("SELECT pg_advisory_xact_lock(hashtextextended(:k, 0))"),
                    {"k": f"tagging_publish:{batch_name}"})
    _check_unused(session, batch_name)
    prefix = _check_folder(s3, batch_name)
    snap = take_snapshot(session, col.id)
    _check_budget(*estimate(session, snap))
    pub = TaggingPublish(batch_name=batch_name, collection_id=col.id, state="queued", s3_prefix=prefix,
                         missing=list(snap["missing"]), snapshot=snap, created_by=staff.id, created_at=now)
    session.add(pub)
    session.flush()
    return pub


# ---------------------------------------------------------------- Label Studio shapes

def _iso(dt: Optional[datetime]) -> Optional[str]:
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%fZ") if dt else None


def ls_frame(index: int) -> int:
    """Label Studio frame number of native 0-based frame `index`: LS frames are 1-based, in the clip's native
    frame space (data.fps is the real decoded fps)."""
    return int(index) + 1


def ls_sequence(track) -> list[dict]:
    seq = []
    for k in track.keyframes:
        x1, y1, x2, y2 = k.xyxy
        seq.append({"frame": ls_frame(k.frame), "x": round(x1 * 100, 4), "y": round(y1 * 100, 4),
                    "width": round((x2 - x1) * 100, 4), "height": round((y2 - y1) * 100, 4),
                    "time": k.t_sec, "enabled": bool(k.enabled), "rotation": 0})
    return seq


def ls_results(tracks, description: Optional[str], frames_count: int, duration: float, origin: str) -> list[dict]:
    """One videorectangle per track and the vlm_description textarea, as Label Studio exports them."""
    out = [{"value": {"framesCount": frames_count, "duration": round(duration, 6),
                      "sequence": ls_sequence(t), "labels": [t.label]},
            "id": t.track_id, "from_name": "bbox", "to_name": "video_full", "type": "videorectangle",
            "origin": origin} for t in tracks]
    if description is not None:
        out.append({"value": {"text": [description]}, "id": "vlm_description", "from_name": "vlm_description",
                    "to_name": "video_crop", "type": "textarea", "origin": origin})
    return out


# ---------------------------------------------------------------- the build

class _Batch:
    """One publish run: copies sources, writes analysis_output files and collects the JSON documents. It writes
    only into its own, new folder (the name is used once), so it never has to skip, verify or delete anything."""

    def __init__(self, session: Session, s3, pub: TaggingPublish, lease):
        self.session, self.s3, self.pub, self.lease = session, s3, pub, lease
        self.prefix = pub.s3_prefix
        self.w = s3.scoped(self.prefix)  # every write of the build goes through this: tagging/<batch>/ only
        self.bucket = s3.bucket
        self.claimed = False  # this run wrote its "building" marker (only then may it mark the batch failed)
        self.final: Optional[str] = None  # the state the database took (ready | partial) before the final marker
        self.marker: dict = {}
        self.missing: list[dict] = []  # clips that failed during the build (they make the batch partial)
        self.ineligible: list[dict] = []  # clips that cannot be published as they are (variable/unknown frame rate)
        self.tasks: list[dict] = []
        self.vlm_lines: list[dict] = []
        self.rows: list[dict] = []  # one per clip, for the workbook and the report
        self.yolo_frames = 0
        self.versions: dict[str, int] = {}
        self.now = datetime.now(timezone.utc)

    def uri(self, rel: str) -> str:
        return f"s3://{self.bucket}/{self.prefix}{rel}"

    def miss(self, event_id: int, reason: str) -> None:
        entry = {"event_id": event_id, "reason": reason}
        if entry not in self.missing:
            self.missing.append(entry)

    def all_missing(self) -> list[dict]:
        return list(self.pub.missing or []) + self.ineligible + self.missing

    # ------------------------------------------------------------ ownership (fencing) and progress

    def counts(self) -> dict:
        return {"tasks": len(self.tasks), "yolo_frames": self.yolo_frames, "vlm_lines": len(self.vlm_lines),
                "missing": self.all_missing()}

    def fence(self) -> None:
        """Raise LeaseLost unless this worker still owns the running publish; records the progress counts and the
        heartbeat in the same statement (its own connection, committed at once)."""
        from . import studio

        m = TaggingPublish
        with self.lease.engine.begin() as conn:
            n = conn.execute(update(m).where(m.id == self.pub.id, m.worker_id == self.lease.worker_id,
                                             m.state == "running")
                             .values(heartbeat_at=func.now(), **self.counts())).rowcount
        if not n:
            self.lease.lost.set()
            raise studio.LeaseLost(self.pub.id)

    def finalize(self, state: str) -> None:
        """running -> `state` in one conditional UPDATE: only while this worker owns the publish, it is running and
        its heartbeat is fresh (not yet sweepable). Runs BEFORE the final marker; LeaseLost when it does not
        match, and then the worker writes nothing more."""
        from . import studio

        m = TaggingPublish
        with self.lease.engine.begin() as conn:
            n = conn.execute(update(m).where(m.id == self.pub.id, m.worker_id == self.lease.worker_id,
                                             m.state == "running",
                                             m.heartbeat_at > func.now() - studio.STALE_RUNNING_AFTER)
                             .values(state=state, error=None, heartbeat_at=func.now(), **self.counts())).rowcount
        if not n:
            self.lease.lost.set()
            raise studio.LeaseLost(self.pub.id)
        self.final = state

    def write_marker(self, **values) -> None:
        self.marker.update(values)
        self.w.put_bytes(self.prefix + MARKER, json.dumps(self.marker, indent=2).encode("utf-8"), "application/json")

    # ------------------------------------------------------------ writes

    def put(self, rel: str, body: bytes, mime: str) -> None:
        self.w.put_bytes(self.prefix + rel, body, mime)

    def upload(self, path: Path, rel: str, mime: str) -> None:
        self.w.upload_file(path, self.prefix + rel, mime)

    # ------------------------------------------------------------ per clip

    def sources(self, ev: Event) -> dict[str, Artifact]:
        return _sources(self.session, ev.id)

    def copy_sources(self, ev: Event, sources: dict[str, Artifact]) -> set[str]:
        """Server-side copies to dataset_multi/<rel>; returns the relative paths copied."""
        done = set()
        for rel, art in sorted(sources.items()):
            try:
                self.w.copy(art.s3_key, self.prefix + f"dataset_multi/{rel}",
                            content_type=_MIME.get(PurePosixPath(rel).suffix.lower()), if_match=art.etag or None)
            except ETagMismatch:
                self.miss(ev.id, f"changed since indexing: {rel}")
                continue
            except Exception as e:  # noqa: BLE001 -- a missing source is recorded, anything else fails the build
                if not _not_found(e):
                    raise
                self.miss(ev.id, f"{SOURCE_MISSING}: {rel}")
                continue
            done.add(rel)
        return done

    def decode(self, ev: Event, clip: Artifact, work: Path, fps: float,
               drop: bool) -> Optional[tuple[int, list[tuple[int, float, Path]]]]:
        """(decoded frame count, every frame as (native index, index / fps, image) -- none for a dropped clip), or
        None when the clip is listed as missing: gone, changed, undecodable, or not constant-rate."""
        try:
            self.s3.download_to(clip.s3_key, work / "src.mp4", if_match=clip.etag or None)
        except ETagMismatch as e:
            self.miss(ev.id, f"frames unavailable: {type(e).__name__}")
            return None
        except Exception as e:  # noqa: BLE001 -- gone since indexing: listed, the job goes on
            if not _not_found(e):
                raise
            self.miss(ev.id, SOURCE_MISSING)
            return None
        try:
            times = labeling.frame_times(work / "src.mp4")
            if not times:
                raise media.MediaError("no frames decoded")
            problem = frame_rate_problem(times, fps)
            if problem:
                self.ineligible.append({"event_id": ev.id, "reason": problem})
                return None
            frames = [] if drop else labeling.extract_frames(work / "src.mp4", work, fps, times=times)
        except media.MediaError as e:
            self.miss(ev.id, f"frames unavailable: {type(e).__name__}")
            return None
        if not drop and len(frames) != len(times):
            self.miss(ev.id, "frames unavailable: frame count")
            return None
        # constant rate (checked above): frame n is at n / fps, exactly the time the analysis pipeline uses
        return len(times), [(index, index / fps, path) for index, _, path in frames]

    def write_frames(self, ev: Event, frames, tracks, camera: str, day: str) -> int:
        """analysis_output/yolo: every frame of the clip (image + label from box_at)."""
        for n, (index, t_sec, path) in enumerate(frames):
            if n and n % FENCE_EVERY == 0:
                self.fence()
            stem = f"{ev.stem}_f{index:04d}"
            self.upload(path, f"analysis_output/yolo/images/{camera}/{day}/{stem}.jpg", "image/jpeg")
            self.put(f"analysis_output/yolo/labels/{camera}/{day}/{stem}.txt",
                     labeling.label_text(labeling.yolo_rows(tracks, t_sec)).encode("utf-8"), "text/plain")
        return len(frames)

    def clip(self, info: dict) -> None:
        self.fence()  # never write a clip for a publish this worker no longer owns
        ev = self.session.get(Event, info["id"])
        if ev is None:
            self.miss(info["id"], "event deleted")
            return
        consent = self.session.scalar(select(Customer.consent_training).join(Device, Device.customer_id == Customer.id)
                                      .where(Device.id == ev.device_pk))
        if not consent:  # withdrawn since the request
            self.miss(ev.id, NO_CONSENT)
            return
        row = labeling.version_row(self.session, ev.id, info["version"])
        if row is None:
            self.miss(ev.id, "annotation missing")
            return
        sources = self.sources(ev)
        clip_rel = next((rel for rel, a in sorted(sources.items()) if a.role == "original_video"), None)
        if clip_rel is None:
            self.miss(ev.id, "clip missing")
            return
        meta = labeling.meta_body(self.session, ev.id)
        fps, _, _ = labeling.clip_timing(ev, meta)
        if frame_rate_problem([], fps):  # never a default rate
            self.ineligible.append({"event_id": ev.id, "reason": RATE_UNKNOWN})
            return
        drop = bool(info.get("drop_clip") or row.drop_clip)
        with tempfile.TemporaryDirectory(prefix="hgtag_") as tmp:
            decoded = self.decode(ev, sources[clip_rel], Path(tmp), fps, drop)
            if decoded is None:
                return
            frames_count, frames = decoded
            self.fence()  # decoding takes a while: check again before the clip's first write
            copied = self.copy_sources(ev, sources)
            if clip_rel not in copied:
                return
            self.task(ev, info, row, meta, sources, clip_rel, copied, fps, frames_count, frames, drop)

    def task(self, ev: Event, info: dict, row, meta: dict, sources: dict[str, Artifact], clip_rel: str,
             copied: set[str], fps: float, frames_count: int, frames, drop: bool) -> None:
        """The clip's Label Studio task, its workbook row and (kept clips) its frames and VLM line."""
        self.versions[str(ev.id)] = row.version
        duration = frames_count / fps  # the decoded clip's length
        camera = ev.camera
        day = ev.day or datetime.fromtimestamp(ev.start_ts, timezone.utc).strftime("%Y-%m-%d")
        clip_id = ev.stem
        tracks = labeling.to_tracks(row.tracks)
        crop_rel = next((rel for rel in sorted(copied) if sources[rel].role == "crop_video"), None)
        video_url = self.uri(f"dataset_multi/{clip_rel}")
        crop_url = self.uri(f"dataset_multi/{crop_rel}") if crop_rel else video_url
        start_local = meta.get("clip_start_local") if isinstance(meta.get("clip_start_local"), str) else "?"
        end_local = meta.get("clip_end_local") if isinstance(meta.get("clip_end_local"), str) else "?"
        kind = meta.get("kind") if isinstance(meta.get("kind"), str) else ev.kind
        main = meta.get("main_stream") if isinstance(meta.get("main_stream"), dict) else {}
        crop = meta.get("vlm_crop") if isinstance(meta.get("vlm_crop"), dict) else {}
        task_id = len(self.tasks) + 1
        data = {
            "video_url": video_url, "vlm_crop_url": crop_url, "fps": fps, "frame_space": FRAME_SPACE,
            "frames_count": frames_count, "vlm_fps": main.get("store_fps", crop.get("fps", fps)),
            "header_info": f"{camera} | {kind} | {start_local} - {end_local.split(' ')[-1]}",
            "vlm_crop_header": "VLM Crop (high-res)" if crop_rel else "No VLM Crop — showing full frame",
            "has_vlm_crop": bool(crop_rel), "camera_name": camera, "kind": kind,
            "clip_start_local": start_local, "clip_end_local": end_local,
            "duration_sec": round(duration, 6), "meta_path": f"meta/{camera}/{day}/{clip_id}.meta.json",
            "clip_id": clip_id, "date": day,
        }
        text_value = DELETE_MARKER if drop else (row.description or "")
        author = self.session.get(Staff, row.author_id) if row.author_id is not None else None
        first = self.session.scalar(select(func.min(Annotation.created_at)).where(Annotation.event_id == ev.id))
        lead = max(0.0, (row.created_at - first).total_seconds()) if first and row.created_at else 0.0
        result = ls_results(tracks, text_value, frames_count, duration, "manual")
        annotation = {
            "id": task_id, "completed_by": row.author_id if row.author_id is not None else 0,
            "created_username": f"{author.email}, {author.id}" if author is not None else (row.author_name or ""),
            "result": result, "was_cancelled": False, "ground_truth": info.get("status") == "reviewed",
            "created_at": _iso(first or row.created_at), "updated_at": _iso(row.created_at), "lead_time": lead,
            "result_count": len(result), "task": task_id,
        }
        suggestions = labeling.suggestions(self.session, self.s3, ev, fps, self.now, store=False)
        run = labeling.guard_run(self.session, ev.id)
        predictions = []
        pred_result = ls_results(suggestions, ev.summary or None, frames_count, duration, "prediction")
        if pred_result:
            predictions.append({"id": task_id, "model_version": labeling.SUGGESTION_MODEL, "score": None,
                                "result": pred_result, "task": task_id})
        self.tasks.append({
            "id": task_id, "data": data, "annotations": [annotation], "predictions": predictions,
            "meta": {"admin_center": {"event_id": ev.id, "annotation_version": row.version,
                                      "status": info.get("status"), "drop_clip": drop,
                                      "description": row.description or ""}},
            "created_at": _iso(first or row.created_at), "updated_at": _iso(row.created_at), "inner_id": task_id,
            "total_annotations": 1, "cancelled_annotations": 0, "total_predictions": len(predictions),
        })
        persons = sum(1 for t in tracks if t.label == "person")
        cars = sum(1 for t in tracks if t.label == "car")
        verdicts = [v for v in (ev.owner_verdicts or []) if isinstance(v, str)]
        record = {
            "task_id": task_id, "event_id": ev.id, "camera_name": camera, "kind": kind, "date": day,
            "clip_id": clip_id, "duration_sec": round(float(ev.duration_sec or 0), 2), "video_s3_path": video_url,
            "vlm_crop_s3_path": crop_url, "num_persons": persons, "num_cars": cars,
            "num_other": len(tracks) - persons - cars, "tracks": len(tracks),
            "keyframes": sum(len(t.keyframes) for t in tracks), "labels": [t.label for t in tracks],
            "description": row.description or "", "ai_description": row.ai_description or "",
            "ai_model": run.model if run is not None else None,
            "ai_prompt_version": run.prompt_version if run is not None else None, "owner_verdicts": verdicts,
            "annotation_version": row.version, "labeled_by": row.author_name or "", "status": info.get("status"),
            "drop_clip": drop, "yolo_frames": 0,
        }
        if not drop:
            record["yolo_frames"] = self.write_frames(ev, frames, tracks, camera, day)
            self.yolo_frames += record["yolo_frames"]
            if record["description"].strip():
                self.vlm_lines.append({
                    "video_s3_path": video_url, "vlm_crop_s3_path": crop_url, "description": record["description"],
                    "camera_name": camera, "duration_sec": record["duration_sec"], "num_persons": persons,
                    "num_cars": cars, "kind": kind, "date": day, "clip_id": clip_id,
                    "ai_description": record["ai_description"], "ai_model": record["ai_model"],
                    "ai_prompt_version": record["ai_prompt_version"], "owner_verdicts": verdicts,
                    "annotation_version": row.version, "labeled_by": record["labeled_by"],
                })
        self.rows.append(record)

    # ------------------------------------------------------------ the batch documents

    def data_yaml(self) -> str:
        names = "\n".join(f"  {i}: {n}" for i, n in enumerate(labeling_names()))
        return f"path: .\ntrain: images\nval: images\nnames:\n{names}\nnc: 9\n"

    def report(self, creator: str) -> str:
        kept = [r for r in self.rows if not r["drop_clip"]]
        dropped = [r for r in self.rows if r["drop_clip"]]
        lines = ["# Label Export Analysis Report", "",
                 f"Admin Center batch `{self.pub.batch_name}`, published {self.now:%Y-%m-%d %H:%M} UTC by {creator}.",
                 "", "## Overview", "",
                 f"- **Total tasks**: {len(self.rows)}", f"- **Kept tasks**: {len(kept)}",
                 f"- **Dropped clips**: {len(dropped)}",
                 f"- **Total bounding-box tracks**: {sum(r['tracks'] for r in kept)}",
                 f"- **YOLO frames (every frame of each kept clip)**: {self.yolo_frames}",
                 f"- **VLM training lines**: {len(self.vlm_lines)}"]
        durations = [r["duration_sec"] for r in kept]
        if durations:
            lines.append(f"- **Duration range**: {min(durations):.1f}s - {max(durations):.1f}s "
                         f"(mean {statistics.mean(durations):.1f}s)")
        lines += ["", "## Per-Camera Breakdown", ""]
        by_cam: dict[str, list[dict]] = defaultdict(list)
        for r in kept:
            by_cam[r["camera_name"]].append(r)
        for cam in sorted(by_cam):
            rs = by_cam[cam]
            lines += [f"### {cam}", "", f"- **Tasks**: {len(rs)}",
                      f"- **Tracks**: {sum(r['tracks'] for r in rs)} (person: {sum(r['num_persons'] for r in rs)}, "
                      f"car: {sum(r['num_cars'] for r in rs)}, other: {sum(r['num_other'] for r in rs)})",
                      f"- **Kinds**: {dict(Counter(r['kind'] for r in rs))}", ""]
        lines += ["## Label Distribution", ""]
        for label, count in Counter(lab for r in kept for lab in r["labels"]).most_common():
            lines.append(f"- **{label}**: {count}")
        lines += ["", "## Labelers", ""]
        for name, count in sorted(Counter(r["labeled_by"] or "unknown" for r in self.rows).items()):
            lines.append(f"- **{name}**: {count} task(s)")
        lines.append("")
        if dropped:
            lines += [f"## Dropped Clips ({DELETE_MARKER} marker)", "", f"- **Total dropped**: {len(dropped)}", "",
                      "| Task ID | Camera | Clip ID |", "|---------|--------|---------|"]
            lines += [f"| {r['task_id']} | {r['camera_name']} | {r['clip_id']} |" for r in dropped]
            lines.append("")
        if self.all_missing():
            lines += ["## Not Published", ""]
            lines += [f"- event {m['event_id']}: {m['reason']}" for m in self.all_missing()]
            lines.append("")
        return "\n".join(lines) + "\n"

    def workbook(self) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Tasks"
        headers = ["task_id", "event_id", "camera_name", "kind", "date", "clip_id", "duration_sec", "video_s3_path",
                   "vlm_crop_s3_path", "num_person_tracks", "num_car_tracks", "num_other_tracks", "total_tracks",
                   "total_keyframes", "description", "ai_description", "ai_model", "ai_prompt_version",
                   "owner_verdicts", "annotation_version", "labeled_by", "status", "drop_clip", "yolo_frames"]
        keys = ["task_id", "event_id", "camera_name", "kind", "date", "clip_id", "duration_sec", "video_s3_path",
                "vlm_crop_s3_path", "num_persons", "num_cars", "num_other", "tracks", "keyframes", "description",
                "ai_description", "ai_model", "ai_prompt_version", "owner_verdicts", "annotation_version",
                "labeled_by", "status", "drop_clip", "yolo_frames"]
        bold, fill = Font(bold=True), PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font, cell.fill = bold, fill
        for i, r in enumerate(self.rows, 2):
            for col, k in enumerate(keys, 1):
                v = r[k]
                ws.cell(row=i, column=col, value=", ".join(v) if isinstance(v, list) else v)
        ws.auto_filter.ref = ws.dimensions
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def labeling_names() -> list[str]:
    """The nine class names in contiguous id order (person, bicycle, car, ...)."""
    from .studio import YOLO_NAMES

    return list(YOLO_NAMES)


def _not_found(exc: Exception) -> bool:
    from .studio import _not_found as nf

    return nf(exc)


def _run(batch: _Batch) -> str:
    session, s3, pub = batch.session, batch.s3, batch.pub
    if s3 is None:
        raise RuntimeError("Storage is not configured")
    _check_folder(s3, pub.batch_name)  # again at build time: this attempt starts from an empty folder
    creator = session.get(Staff, pub.created_by)
    creator_name = creator.name if creator is not None else ""
    batch.marker = {"created_by": creator_name, "created_utc": _iso(pub.created_at),
                    "collection_id": pub.collection_id, "publish_id": pub.id, "batch_name": pub.batch_name,
                    "frame_space": FRAME_SPACE}
    batch.fence()
    batch.write_marker(state="building")  # first: claims the folder
    batch.claimed = True
    for info in (pub.snapshot or {}).get("events", []):
        batch.clip(info)
        batch.fence()  # progress counts, and the worker still owns the job
    tasks = batch.tasks
    batch.put(f"{pub.batch_name}.json", json.dumps(tasks, ensure_ascii=False, indent=2).encode("utf-8"),
              "application/json")
    plain = [{k: v for k, v in t.items() if k != "annotations"} for t in tasks]
    batch.put("dataset_multi/label_studio_tasks.json", json.dumps(plain, ensure_ascii=False, indent=2).encode("utf-8"),
              "application/json")
    batch.put("dataset_multi/label_studio_config.xml", label_config_xml().encode("utf-8"), "application/xml")
    batch.put("analysis_output/yolo/data.yaml", batch.data_yaml().encode("utf-8"), "application/yaml")
    batch.put("analysis_output/vlm_training.jsonl",
              "".join(json.dumps(line, ensure_ascii=False) + "\n" for line in batch.vlm_lines).encode("utf-8"),
              "application/jsonl")
    batch.put("analysis_output/summary_report.md", batch.report(creator_name).encode("utf-8"), "text/markdown")
    batch.put("analysis_output/analysis.xlsx", batch.workbook(),
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    state = "partial" if batch.missing else "ready"
    batch.finalize(state)  # the database first (this worker, running, fresh heartbeat), then the marker
    batch.write_marker(state=state, finished_utc=_iso(datetime.now(timezone.utc)), annotations=batch.versions,
                       tasks=len(tasks), yolo_frames=batch.yolo_frames, vlm_lines=len(batch.vlm_lines))
    return state


# ---------------------------------------------------------------- the job (export worker pool, leased)

def claim_publish(session: Session, publish_id: int, worker_id: str) -> bool:
    row = session.execute(text(
        "UPDATE tagging_publishes SET state='running', worker_id=:w, heartbeat_at=now(), error=NULL "
        "WHERE id=:id AND state='queued' RETURNING id"), {"w": worker_id, "id": publish_id}).first()
    session.commit()
    return row is not None


def build_publish(session: Session, s3, publish_id: int, worker_id: Optional[str] = None) -> Optional[TaggingPublish]:
    """queued -> running -> ready | partial | failed, under a heartbeat lease like training exports. The final
    state is set by the run itself (_Batch.finalize) before its final marker."""
    from . import studio

    worker_id = worker_id or studio.new_worker_id()
    if not claim_publish(session, publish_id, worker_id):
        return session.get(TaggingPublish, publish_id, populate_existing=True)
    pub = session.get(TaggingPublish, publish_id, populate_existing=True)
    lease = studio._Lease(session.get_bind(), publish_id, worker_id, model=TaggingPublish)
    lease.start()
    batch = _Batch(session, s3, pub, lease) if s3 is not None else None
    try:
        if batch is None:
            raise RuntimeError("Storage is not configured")
        _run(batch)
    except studio.LeaseLost:  # swept as stale and maybe taken over: stop without writing anything more
        session.rollback()
        log.warning("publish %s: lease lost (swept as stale); stopped", publish_id)
    except Exception as e:  # noqa: BLE001 -- any failure ends the publish as failed, never stuck in running
        log.exception("publish %s failed", publish_id)
        session.rollback()
        counts = batch.counts() if batch is not None else {}
        owned = studio._finish(session, publish_id, worker_id, model=TaggingPublish, state="failed",
                               error=studio.error_text(e), **counts)
        if not owned and batch is not None and batch.final:  # the database said ready, the marker did not land
            n = session.execute(update(TaggingPublish).where(
                TaggingPublish.id == publish_id, TaggingPublish.worker_id == worker_id,
                TaggingPublish.state == batch.final).values(state="failed", error=studio.error_text(e), **counts)
                .execution_options(synchronize_session=False)).rowcount
            session.commit()
            owned = bool(n)
        if owned and batch is not None and batch.claimed:
            try:  # best effort: the name is spent either way, and this folder is never written again
                batch.write_marker(state="failed", finished_utc=_iso(datetime.now(timezone.utc)))
            except Exception:  # noqa: BLE001 -- the batch stays "building", which the analysis refuses as well
                log.exception("publish %s: could not write the failed marker", publish_id)
    finally:
        lease.stop()
    return session.get(TaggingPublish, publish_id, populate_existing=True)


def run_publish_job(sessionmaker, s3, publish_id: int) -> None:
    session = sessionmaker()
    try:
        build_publish(session, s3, publish_id)
    except Exception:  # noqa: BLE001 -- the sweep fails it later if even marking it failed broke
        log.exception("publish job %s crashed", publish_id)
    finally:
        session.close()


def sweep_stale_publishes(session: Session, now: datetime) -> int:
    """Fail publishes whose worker is gone (the training exports' rule)."""
    from . import studio

    beat = func.coalesce(TaggingPublish.heartbeat_at, TaggingPublish.created_at)
    stale = or_((TaggingPublish.state == "running") & (beat < now - studio.STALE_RUNNING_AFTER),
                (TaggingPublish.state == "queued") & (TaggingPublish.created_at < now - studio.STALE_QUEUED_AFTER))
    result = session.execute(update(TaggingPublish).where(stale).values(state="failed", error=studio.STALE_ERROR)
                             .execution_options(synchronize_session=False))
    session.flush()
    return result.rowcount or 0
