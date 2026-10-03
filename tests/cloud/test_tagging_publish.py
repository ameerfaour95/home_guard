"""Publish a Studio collection as a tagging batch: s3 tagging/<batch>/ in the layout of the Label Studio batches,
parsed by the same logic home_guard_project/analysis/analyze.py uses (vendored below), the read-only rule for
existing batches, and writes confined to tagging/<batch>/."""
import io
import json
import os
import re
import statistics
from typing import Any, Dict, List

import pytest
from sqlalchemy import select

from home_guard_project.cloud import models as m
from home_guard_project.cloud import tagging
from home_guard_project.cloud.db import session_scope

from . import builders as b
from .test_annotation_export import FRAMES, PERSON, _annotate, make_clip, needs_ffmpeg
from .test_event_routes import _event_id, s3client  # noqa: F401
from .test_studio import _get, _keys, _make_collection, _seed

BATCH = "night_batch"
P = f"tagging/{BATCH}/"


# ---------------------------------------------------------------- vendored from the founder's pipeline
# Minimal copies of home_guard_project/analysis/analyze.py (parse_export, _is_delete_marker),
# home_guard_project/analysis/utils/s3.py (meta_path_to_*) and home_guard_project/analysis/utils/ls_convert.py
# (interpolate_keyframes) at branch beelink-collector-box, so this test proves a published <batch>.json parses with
# exactly that logic. The only change: parse_export takes the loaded list instead of a path and config, and keeps
# the fields the test checks.

def meta_path_to_clip_id(meta_path: str) -> str:
    basename = os.path.basename(meta_path)
    return re.sub(r"\.meta\.json$", "", basename)


def meta_path_to_camera_date(meta_path: str) -> tuple:
    parts = meta_path.replace("\\", "/").split("/")
    camera = parts[1] if len(parts) > 1 else "unknown"
    date = parts[2] if len(parts) > 2 else "unknown"
    return camera, date


def meta_path_to_s3_clip(meta_path: str, bucket: str, prefix: str) -> str:
    rel = meta_path.replace("\\", "/")
    rel = re.sub(r"^meta/", "clips/", rel)
    rel = re.sub(r"\.meta\.json$", ".mp4", rel)
    return f"s3://{bucket}/{prefix}/{rel}"


def parse_export(raw_tasks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    parsed = []
    for raw in raw_tasks:
        data = raw.get("data", {})
        meta_path = data.get("meta_path", "")
        camera, date = meta_path_to_camera_date(meta_path)
        pt = {"task_id": raw.get("id", 0), "camera_name": data.get("camera_name", camera),
              "kind": data.get("kind", "unknown"), "date": date,
              "duration_sec": float(data.get("duration_sec", 0)), "fps": float(data.get("fps", 7)),
              "clip_id": meta_path_to_clip_id(meta_path), "meta_path": meta_path,
              "s3_clip_url": meta_path_to_s3_clip(meta_path, b.BUCKET, f"{P}dataset_multi".rstrip("/")),
              "tracks": [], "vlm_description": "", "annotator_id": None, "lead_time_sec": 0.0}
        annotations = raw.get("annotations", [])
        annotator_ids = set()
        lead_times = []
        for ann in annotations:
            if ann.get("was_cancelled"):
                continue
            cby = ann.get("completed_by")
            if cby is not None:
                annotator_ids.add(cby)
            lt = ann.get("lead_time")
            if lt is not None:
                lead_times.append(float(lt))
            for r in ann.get("result", []):
                rtype = r.get("type")
                val = r.get("value", {})
                if rtype == "videorectangle":
                    labels = val.get("labels", [])
                    seq = val.get("sequence", [])
                    frames_count = int(val.get("framesCount", 0))
                    duration = float(val.get("duration", 0))
                    enabled_kfs = [kf for kf in seq if kf.get("enabled", True)]
                    frame_nums = [int(kf.get("frame", 0)) for kf in enabled_kfs]
                    areas = [kf.get("width", 0) * kf.get("height", 0) for kf in enabled_kfs]
                    pt["tracks"].append({
                        "label": labels[0] if labels else "unknown", "num_keyframes": len(seq),
                        "first_frame": min(frame_nums) if frame_nums else 0,
                        "last_frame": max(frame_nums) if frame_nums else 0, "frames_count": frames_count,
                        "duration": duration, "avg_box_area_pct": statistics.mean(areas) if areas else 0,
                        "has_enabled_false": any(not kf.get("enabled", True) for kf in seq), "sequence": seq})
                elif rtype == "textarea" and r.get("from_name") == "vlm_description":
                    texts = val.get("text", [])
                    if texts:
                        pt["vlm_description"] = texts[0]
        if annotator_ids:
            pt["annotator_id"] = min(annotator_ids)
        if lead_times:
            pt["lead_time_sec"] = sum(lead_times) / len(lead_times)
        parsed.append(pt)
    return parsed


def is_delete_marker(description: str) -> bool:
    return description.strip().lower() == "[delete]"


def interpolate_keyframes(sequence, frames_count):
    if not sequence:
        return {}
    sorted_kfs = sorted(sequence, key=lambda kf: kf.get("frame", 0))
    result = {}
    lerp = lambda a, b_, t: a + (b_ - a) * t  # noqa: E731
    for i, kf in enumerate(sorted_kfs):
        if not kf.get("enabled", True):
            continue
        frame = int(kf["frame"])
        end_frame = int(sorted_kfs[i + 1]["frame"]) if i + 1 < len(sorted_kfs) else frames_count
        for f in range(frame, end_frame + 1):
            if f > frames_count:
                break
            t = 0.0 if frame == end_frame else (f - frame) / (end_frame - frame)
            if i + 1 < len(sorted_kfs):
                nxt = sorted_kfs[i + 1]
                if not nxt.get("enabled", True):
                    if f >= end_frame:
                        break
                    box = {k: kf[k] for k in ("x", "y", "width", "height")}
                else:
                    box = {k: lerp(kf[k], nxt[k], t) for k in ("x", "y", "width", "height")}
            else:
                box = {k: kf[k] for k in ("x", "y", "width", "height")}
            result[f] = box
    return result


# labeling/tasks.py _LABEL_CONFIG_XML filled with the nine labels, as tagging/ameer_house_batch_2/dataset_multi/
# label_studio_config.xml holds it
EXPECTED_XML = """<View>
  <Header value="$header_info"/>
  <View style="display:flex; gap:16px;">
    <View style="flex:1;">
      <Header value="Full Frame — YOLO Re-tagging"/>
      <Labels name="label" toName="video_full">
    <Label value="bicycle"/>
    <Label value="bird"/>
    <Label value="bus"/>
    <Label value="car"/>
    <Label value="cat"/>
    <Label value="dog"/>
    <Label value="motorcycle"/>
    <Label value="person"/>
    <Label value="truck"/>
      </Labels>
      <Video name="video_full" value="$video_url" frameRate="$fps"/>
      <VideoRectangle name="bbox" toName="video_full"/>
    </View>
    <View style="flex:1;">
      <Header value="$vlm_crop_header"/>
      <Video name="video_crop" value="$vlm_crop_url" frameRate="$vlm_fps"/>
      <Header value="Scene Description (for VLM training)"/>
      <TextArea name="vlm_description" toName="video_crop"
                rows="4" editable="true"
                placeholder="Describe what happens in the video..."/>
    </View>
  </View>
</View>
"""


# ---------------------------------------------------------------- fixtures

@pytest.fixture()
def labeled(client, s3client, staff_factory, tmp_path):
    clip = make_clip(tmp_path)
    _seed(client, s3client, overrides={b.COLLECT_CLIP: clip, b.TRAIN_CLIP: clip, b.PAUSED_CLIP: clip})
    labeler, _, _, lab = staff_factory("labeler")
    admin, _, _, adm = staff_factory("admin")
    collect, alert, paused, fp = (_event_id(client, s) for s in (b.COLLECT_STEM, b.STEM, b.PAUSED_STEM, b.FP_STEM))
    _annotate(client, lab, collect, tracks=[PERSON], description="One person crosses the yard.")
    _annotate(client, lab, alert, description="A courier leaves a parcel.")
    assert client.post(f"/v1/events/{alert}/annotation/review", headers=adm,
                       json={"decision": "accept", "version": 1}).status_code == 200
    _annotate(client, lab, paused, description="nothing here", drop_clip=True)
    _annotate(client, lab, fp, description="draft only", status="edited")
    cid = _make_collection(client, adm, [collect, alert, paused, fp], name="batch three")
    return {"adm": adm, "admin": admin, "cid": cid, "collect": collect, "alert": alert, "paused": paused, "fp": fp,
            "labeler": labeler, "lab": lab}


def _publish(client, h, cid, name=BATCH):
    return client.post(f"/v1/studio/collections/{cid}/publish", headers=h, json={"batch_name": name})


def _all_keys(s3client):
    return set(_keys(s3client, ""))


# ---------------------------------------------------------------- the layout

@needs_ffmpeg
def test_publish_writes_the_label_studio_batch_layout(client, s3client, labeled):
    before = _all_keys(s3client)
    r = _publish(client, labeled["adm"], labeled["cid"])
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["state"] == "ready", out
    assert (out["batch_name"], out["s3_prefix"], out["tasks"], out["yolo_frames"], out["vlm_lines"]) == (
        BATCH, P, 3, 2 * FRAMES, 2)
    assert out["missing"] == [{"event_id": labeled["fp"], "reason": "not labeled"}]
    assert out["created_by"] == labeled["admin"].name
    written = _all_keys(s3client) - before
    assert written and all(k.startswith(P) for k in written)  # nothing outside tagging/<batch>/

    # dataset_multi: server-side copies under the box's relative paths, untouched
    day, stem = "back_door/2026-10-02", b.COLLECT_STEM
    assert _get(s3client, f"{P}dataset_multi/clips/{day}/{stem}.mp4") == _get(s3client, b.COLLECT_CLIP)
    assert _get(s3client, f"{P}dataset_multi/meta/{day}/{stem}.meta.json") == _get(s3client, b.COLLECT_META)
    assert _get(s3client, f"{P}dataset_multi/yolo/labels/{day}/{stem}_f0000.txt") == _get(s3client, b.YOLO_LABELS[0])
    assert f"{P}dataset_multi/yolo/images/{day}/{stem}_f0002.jpg" in written
    assert f"{P}dataset_multi/responses/front_side/2026-10-03/{b.STEM}.model_raw.txt" in written
    assert _get(s3client, f"{P}dataset_multi/label_studio_config.xml").decode("utf-8") == EXPECTED_XML

    # analysis_output/yolo: every frame, native 0-based index, labels from box_at, contiguous ids
    images = sorted(k for k in written if k.startswith(f"{P}analysis_output/yolo/images/{day}/"))
    assert [k.rsplit("/", 1)[1] for k in images] == [f"{stem}_f{i:04d}.jpg" for i in range(FRAMES)]
    label = lambda i: _get(s3client, f"{P}analysis_output/yolo/labels/{day}/{stem}_f{i:04d}.txt").decode()  # noqa
    assert label(0) == "0 0.200000 0.300000 0.200000 0.400000\n"
    assert label(3) == "0 0.371429 0.300000 0.200000 0.400000\n" and label(9) == ""
    alert_labels = [k for k in written if k.startswith(f"{P}analysis_output/yolo/labels/front_side/2026-10-03/")]
    assert len(alert_labels) == FRAMES and all(_get(s3client, k) == b"" for k in alert_labels)  # human negatives
    assert not [k for k in written if "/left_side_1/" in k and "/analysis_output/" in k]  # dropped clip
    assert _get(s3client, f"{P}analysis_output/yolo/data.yaml").decode() == (
        "path: .\ntrain: images\nval: images\nnames:\n  0: person\n  1: bicycle\n  2: car\n  3: motorcycle\n"
        "  4: bus\n  5: truck\n  6: bird\n  7: cat\n  8: dog\nnc: 9\n")

    # vlm_training.jsonl: the existing fields plus the provenance
    lines = [json.loads(x) for x in _get(s3client, f"{P}analysis_output/vlm_training.jsonl").decode().splitlines()]
    by_clip = {line["clip_id"]: line for line in lines}
    assert set(by_clip) == {b.COLLECT_STEM, b.STEM}
    line = by_clip[b.STEM]
    assert {"video_s3_path", "vlm_crop_s3_path", "description", "camera_name", "duration_sec", "num_persons",
            "num_cars", "kind", "date", "clip_id", "ai_description", "ai_model", "ai_prompt_version",
            "owner_verdicts", "annotation_version", "labeled_by"} == set(line)
    assert line["description"] == "A courier leaves a parcel." and line["labeled_by"] == labeled["labeler"].name
    assert line["video_s3_path"] == f"s3://{b.BUCKET}/{P}dataset_multi/clips/front_side/2026-10-03/{b.STEM}.mp4"
    assert line["ai_description"].startswith("A person appears") and line["owner_verdicts"] == ["true_alert"]
    report = _get(s3client, f"{P}analysis_output/summary_report.md").decode()
    assert "## Per-Camera Breakdown" in report and "Dropped Clips" in report and labeled["labeler"].name in report
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(_get(s3client, f"{P}analysis_output/analysis.xlsx")))["Tasks"]
    assert ws.max_row == 4 and ws.cell(row=1, column=1).value == "task_id"

    marker = json.loads(_get(s3client, P + "_admin_center.json"))
    assert marker["collection_id"] == labeled["cid"] and marker["created_by"] == labeled["admin"].name
    assert marker["annotations"] == {str(labeled["collect"]): 1, str(labeled["alert"]): 1,
                                     str(labeled["paused"]): 1}
    plain = json.loads(_get(s3client, f"{P}dataset_multi/label_studio_tasks.json"))
    assert len(plain) == 3 and all("annotations" not in t and t["data"]["meta_path"] for t in plain)


@needs_ffmpeg
def test_batch_json_parses_with_the_analysis_pipeline(client, s3client, labeled):
    assert _publish(client, labeled["adm"], labeled["cid"]).json()["state"] == "ready"
    raw = json.loads(_get(s3client, f"{P}{BATCH}.json"))
    tasks = {t["clip_id"]: t for t in parse_export(raw)}
    assert set(tasks) == {b.COLLECT_STEM, b.STEM, b.PAUSED_STEM}
    collect = tasks[b.COLLECT_STEM]
    assert (collect["camera_name"], collect["date"], collect["kind"], collect["fps"]) == (
        "back_door", "2026-10-02", "trigger", 7.0)
    assert collect["annotator_id"] == labeled["labeler"].id
    assert collect["vlm_description"] == "One person crosses the yard."
    assert collect["s3_clip_url"] == f"s3://{b.BUCKET}/{P}dataset_multi/clips/back_door/2026-10-02/{b.COLLECT_STEM}.mp4"
    [track] = collect["tracks"]
    assert track["label"] == "person" and track["num_keyframes"] == 3 and track["has_enabled_false"]
    assert [(k["frame"], k["enabled"]) for k in track["sequence"]] == [(1, True), (8, True), (10, False)]
    assert track["sequence"][0] | {} == {"frame": 1, "x": 10.0, "y": 10.0, "width": 20.0, "height": 40.0,
                                         "time": 0.0, "enabled": True, "rotation": 0}
    assert track["frames_count"] == 63 and track["duration"] == pytest.approx(9.0)
    boxes = interpolate_keyframes(track["sequence"], track["frames_count"])
    assert boxes[1]["x"] == 10.0 and boxes[8]["x"] == pytest.approx(50.0) and 10 not in boxes
    deleted = [t for t in tasks.values() if is_delete_marker(t["vlm_description"])]
    assert [t["clip_id"] for t in deleted] == [b.PAUSED_STEM]  # drop_clip: kept in the JSON, [delete] marker
    assert tasks[b.STEM]["tracks"] == [] and tasks[b.STEM]["vlm_description"] == "A courier leaves a parcel."
    task = next(t for t in raw if t["data"]["clip_id"] == b.COLLECT_STEM)
    assert task["predictions"][0]["model_version"] == "yolov8n-weak-labels"
    assert {r["type"] for r in task["predictions"][0]["result"]} == {"videorectangle"}


# ---------------------------------------------------------------- the read-only rule

@needs_ffmpeg
def test_existing_batches_are_never_touched(client, s3client, labeled):
    adm, cid = labeled["adm"], labeled["cid"]
    b.put(s3client, "tagging/ameer_house_batch_2/ameer_house_batch_2.json", "[]")
    b.put(s3client, "tagging/someone_else/someone_else.json", "[]")
    before = {k: _get(s3client, k) for k in _all_keys(s3client)}
    for name in ("ameer_house_batch_1", "ameer_house_batch_2", "uca_dataset_batch", "smarthome_dataset_batch",
                 "someone_else"):
        r = _publish(client, adm, cid, name)
        assert r.status_code == 409, (name, r.text)
    assert {k: _get(s3client, k) for k in _all_keys(s3client)} == before
    assert _publish(client, adm, cid, "Bad-Name").status_code == 422


def test_publish_is_admin_only(client, s3client, staff_factory):
    _seed(client, s3client)
    _, _, _, adm = staff_factory("admin")
    cid = _make_collection(client, adm, [_event_id(client, b.COLLECT_STEM)])
    for role in ("labeler", "support"):
        _, _, _, h = staff_factory(role)
        assert _publish(client, h, cid).status_code == 403
        assert client.get("/v1/studio/publishes", headers=h).status_code == 403
    assert _publish(client, adm, 999999).status_code == 404


def test_nothing_labeled_publishes_an_empty_batch_with_everything_missing(client, s3client, staff_factory):
    _seed(client, s3client)
    _, _, _, adm = staff_factory("admin")
    eid = _event_id(client, b.COLLECT_STEM)
    cid = _make_collection(client, adm, [eid])
    out = _publish(client, adm, cid, "empty_one").json()
    assert (out["state"], out["tasks"], out["missing"]) == ("ready", 0, [{"event_id": eid, "reason": "not labeled"}])
    assert json.loads(_get(s3client, "tagging/empty_one/empty_one.json")) == []
    with session_scope(client.app.state.engine) as s:
        assert s.scalar(select(m.AuditLog.action).where(m.AuditLog.action == "tagging_publish")) == "tagging_publish"


def test_sweep_fails_stale_publishes(client, staff_factory):
    from datetime import datetime, timedelta, timezone

    staff, _, _, _ = staff_factory("admin")
    now = datetime(2026, 10, 3, 12, 0, tzinfo=timezone.utc)
    with session_scope(client.app.state.engine) as s:
        s.add(m.TaggingPublish(batch_name="x", state="running", created_by=staff.id,
                               created_at=now - timedelta(minutes=10), heartbeat_at=now - timedelta(minutes=6)))
        s.add(m.TaggingPublish(batch_name="y", state="running", created_by=staff.id,
                               created_at=now - timedelta(minutes=10), heartbeat_at=now - timedelta(minutes=1)))
        s.flush()
        assert tagging.sweep_stale_publishes(s, now) == 1


def test_app_startup_sweeps_stale_publishes(client, staff_factory):
    from datetime import timedelta

    from home_guard_project.cloud.app import sweep_exports

    staff, _, _, _ = staff_factory("admin")
    now = client.app.state.clock()
    with session_scope(client.app.state.engine) as s:
        s.add(m.TaggingPublish(batch_name="z", state="queued", created_by=staff.id,
                               created_at=now - timedelta(hours=1)))
    sweep_exports(client.app)
    with session_scope(client.app.state.engine) as s:
        assert s.scalar(select(m.TaggingPublish.state).where(m.TaggingPublish.batch_name == "z")) == "failed"


@needs_ffmpeg
def test_a_crash_marks_the_publish_failed_without_paths(client, s3client, labeled, monkeypatch):
    from home_guard_project.cloud import labeling

    def boom(*a, **k):
        raise RuntimeError("disk full at C:/secret/path/clip.mp4")

    monkeypatch.setattr(labeling, "extract_frames", boom)
    out = _publish(client, labeled["adm"], labeled["cid"]).json()
    assert out["state"] == "failed"
    with session_scope(client.app.state.engine) as s:
        error = s.scalar(select(m.TaggingPublish.error))
    assert error.startswith("RuntimeError") and "secret" not in error


# ---------------------------------------------------------------- fix round (fix-label.md D1, D3, D4, D7, D8)

def _labeled_at(client, s3client, staff_factory, tmp_path, rate, frames):
    """The collection clip as a `frames`-frame clip at `rate` fps (event, meta and the video agree)."""
    from home_guard_project.cloud.models import Event, RawRevision

    clip = make_clip(tmp_path, frames=frames, rate=rate)
    _seed(client, s3client, overrides={b.COLLECT_CLIP: clip})
    collect = _event_id(client, b.COLLECT_STEM)
    with session_scope(client.app.state.engine) as s:
        ev = s.get(Event, collect)
        ev.fps, ev.duration_sec = float(rate), frames / rate
        for rev in s.scalars(select(RawRevision).where(RawRevision.s3_key == b.COLLECT_META)):
            rev.body = {**rev.body, "frames_written": frames, "fps_estimated": float(rate)}
    labeler, _, _, lab = staff_factory("labeler")
    _, _, _, adm = staff_factory("admin")
    return collect, lab, adm


@needs_ffmpeg
def test_native_frame_space_at_30_fps(client, s3client, staff_factory, tmp_path):
    # D1: at 30 fps, adjacent keyframes on native frames 0 and 1 stay two LS frames (the old 10 fps cap merged them)
    collect, lab, adm = _labeled_at(client, s3client, staff_factory, tmp_path, rate=30, frames=30)
    track = {"track_id": "x", "label": "person", "source": "human", "keyframes": [
        {"frame": 0, "t_sec": 0.0, "xyxy": [0.1, 0.1, 0.3, 0.5], "enabled": True},
        {"frame": 1, "t_sec": 1 / 30, "xyxy": [0.5, 0.1, 0.7, 0.5], "enabled": True},
        {"frame": 2, "t_sec": 2 / 30, "xyxy": [0.5, 0.1, 0.7, 0.5], "enabled": False}]}
    _annotate(client, lab, collect, tracks=[track], description="A person.")
    cid = _make_collection(client, adm, [collect], name="thirty")
    out = _publish(client, adm, cid, "thirty_fps").json()
    assert out["state"] == "ready", out
    p = "tagging/thirty_fps/"
    [task] = json.loads(_get(s3client, f"{p}thirty_fps.json"))
    assert task["data"]["fps"] == 30.0 and task["data"]["frame_space"] == "native"
    [rect] = [r for r in task["annotations"][0]["result"] if r["type"] == "videorectangle"]
    assert rect["value"]["framesCount"] == 30
    seq = rect["value"]["sequence"]
    assert [(k["frame"], k["enabled"]) for k in seq] == [(1, True), (2, True), (3, False)]
    assert [k["time"] for k in seq] == pytest.approx([0.0, 1 / 30, 2 / 30])
    assert [k["x"] for k in seq[:2]] == [10.0, 50.0]
    assert json.loads(_get(s3client, p + "_admin_center.json"))["frame_space"] == "native"
    day, stem = "back_door/2026-10-02", b.COLLECT_STEM
    keys = _all_keys(s3client)
    images = sorted(k.rsplit("/", 1)[1][:-4] for k in keys if k.startswith(f"{p}analysis_output/yolo/images/{day}/"))
    labels = sorted(k.rsplit("/", 1)[1][:-4] for k in keys if k.startswith(f"{p}analysis_output/yolo/labels/{day}/"))
    assert images == labels == [f"{stem}_f{n:04d}" for n in range(30)]  # every native frame n: image + label
    label = lambda n: _get(s3client, f"{p}analysis_output/yolo/labels/{day}/{stem}_f{n:04d}.txt").decode()  # noqa
    assert label(0).startswith("0 0.200000 ") and label(1).startswith("0 0.600000 ") and label(2) == ""


@needs_ffmpeg
def test_a_published_batch_is_immutable(client, s3client, labeled):
    # D3: a ready batch is never rewritten
    assert _publish(client, labeled["adm"], labeled["cid"]).json()["state"] == "ready"
    before = {k: _get(s3client, k) for k in _all_keys(s3client)}
    r = _publish(client, labeled["adm"], labeled["cid"])
    assert r.status_code == 409 and r.json()["detail"] == "This batch is published; choose a new name"
    assert {k: _get(s3client, k) for k in _all_keys(s3client)} == before


class _Writes:
    """Records every key the S3 wrapper writes (put, upload, copy, delete)."""

    def __init__(self, monkeypatch):
        from home_guard_project.cloud.s3 import S3

        self.keys: list[str] = []
        self.deleted: list[str] = []
        for name, pos in (("put_bytes", 0), ("upload_file", 1), ("copy", 1)):
            real = getattr(S3, name)

            def wrapper(this, *a, _real=real, _pos=pos, **k):
                self.keys.append(a[_pos])
                return _real(this, *a, **k)

            monkeypatch.setattr(S3, name, wrapper)
        real_delete = S3.delete_keys

        def delete(this, keys, _real=real_delete):
            self.deleted += list(keys)
            return _real(this, keys)

        monkeypatch.setattr(S3, "delete_keys", delete)
        monkeypatch.setattr(S3, "delete_prefix", lambda this, prefix: self.deleted.append(prefix) or 0)


@needs_ffmpeg
def test_a_failed_publish_resumes_without_rewriting_or_deleting(client, s3client, labeled, monkeypatch):
    # D3: failed -> resumed: objects already there with the recorded hash are skipped, nothing is deleted, and the
    # result equals a clean publish
    from home_guard_project.cloud import labeling

    adm, cid = labeled["adm"], labeled["cid"]
    real = labeling.extract_frames
    calls = []

    def second_fails(*a, **k):
        calls.append(1)
        if len(calls) == 2:
            raise RuntimeError("disk full")
        return real(*a, **k)

    monkeypatch.setattr(labeling, "extract_frames", second_fails)
    assert _publish(client, adm, cid).json()["state"] == "failed"
    marker = json.loads(_get(s3client, P + "_admin_center.json"))
    assert marker["state"] == "failed"
    first_frames = {k for k in _all_keys(s3client) if k.startswith(P + "analysis_output/yolo/")}
    assert len(first_frames) == 2 * FRAMES  # the first clip's images and labels were written
    assert {P + rel for rel in marker["objects"]} >= first_frames  # with their sha256
    b.put(s3client, P + "notes/keep_me.txt", "an earlier file")
    monkeypatch.setattr(labeling, "extract_frames", real)
    writes = _Writes(monkeypatch)
    out = _publish(client, adm, cid).json()
    assert out["state"] == "ready", out
    assert not (set(writes.keys) & first_frames), sorted(set(writes.keys) & first_frames)[:3]
    assert writes.deleted == [] and P + "notes/keep_me.txt" in _all_keys(s3client)
    assert json.loads(_get(s3client, P + "_admin_center.json"))["state"] == "ready"
    assert _publish(client, adm, cid, "clean_copy").json()["state"] == "ready"

    def frames_of(prefix):
        return {k[len(prefix):]: _get(s3client, k) for k in _all_keys(s3client)
                if k.startswith(prefix + "analysis_output/yolo/")}

    assert frames_of(P) == frames_of("tagging/clean_copy/")


@needs_ffmpeg
def test_an_unfinished_admin_center_batch_is_resumed_and_nothing_deleted(client, s3client, labeled):
    adm, cid = labeled["adm"], labeled["cid"]
    b.put(s3client, P + "_admin_center.json", {"created_by": "earlier", "state": "partial"})
    b.put(s3client, P + "analysis_output/yolo/labels/old/2026-01-01/gone_f0000.txt", "0 0.5 0.5 0.1 0.1\n")
    r = _publish(client, adm, cid)
    assert r.status_code == 200 and r.json()["state"] == "ready", r.text
    assert P + "analysis_output/yolo/labels/old/2026-01-01/gone_f0000.txt" in _all_keys(s3client)  # never deleted
    assert json.loads(_get(s3client, P + "_admin_center.json"))["state"] == "ready"


@needs_ffmpeg
def test_a_swept_worker_never_writes_after_a_new_worker_took_over(client, s3client, labeled, monkeypatch):
    # D4: worker A misses its heartbeats and is swept; B publishes the same batch to the end; A wakes up inside its
    # run and must stop without writing anything
    from sqlalchemy import update

    engine, adm, cid = client.app.state.engine, labeled["adm"], labeled["cid"]
    client.app.state.export_runner = lambda job: None  # A stays queued
    assert _publish(client, adm, cid).json()["state"] == "queued"
    with session_scope(engine) as s:
        a_id = s.scalar(select(m.TaggingPublish.id).where(m.TaggingPublish.batch_name == BATCH))
    real_clip = tagging._Batch.clip
    seen = {"a": 0}

    def clip(self, info):
        if self.pub.id == a_id:
            seen["a"] += 1
            if seen["a"] == 2:
                with session_scope(engine) as s:
                    s.execute(update(m.TaggingPublish).where(m.TaggingPublish.id == a_id)
                              .values(state="failed", error="stale"))
                client.app.state.export_runner = lambda job: job()
                assert _publish(client, adm, cid).json()["state"] == "ready"
                seen["after_b"] = {k: _get(s3client, k) for k in _all_keys(s3client)}
        return real_clip(self, info)

    monkeypatch.setattr(tagging._Batch, "clip", clip)
    session = client.app.state.sessionmaker()
    try:
        tagging.build_publish(session, client.app.state.s3, a_id, worker_id="worker-a")
    finally:
        session.close()
    assert "after_b" in seen
    assert {k: _get(s3client, k) for k in _all_keys(s3client)} == seen["after_b"]  # A wrote nothing after B
    marker = json.loads(_get(s3client, P + "_admin_center.json"))
    assert marker["state"] == "ready" and marker["publish_id"] != a_id
    with session_scope(engine) as s:
        assert s.get(m.TaggingPublish, a_id).state == "failed"


@needs_ffmpeg
def test_a_clip_missing_at_download_is_listed_and_the_rest_published(client, s3client, labeled, monkeypatch):
    # D7: the clip vanished between the copy and the frame download: listed in missing, the job goes on
    from botocore.exceptions import ClientError

    from home_guard_project.cloud.s3 import S3

    real = S3.download_to

    def download(this, key, path, if_match=None, **k):
        if key == b.COLLECT_CLIP:
            raise ClientError({"Error": {"Code": "NoSuchKey", "Message": "gone"},
                               "ResponseMetadata": {"HTTPStatusCode": 404}}, "GetObject")
        return real(this, key, path, if_match=if_match, **k)

    monkeypatch.setattr(S3, "download_to", download)
    out = _publish(client, labeled["adm"], labeled["cid"]).json()
    assert out["state"] == "partial", out
    assert {"event_id": labeled["collect"], "reason": "source missing"} in out["missing"]
    assert out["tasks"] == 3 and out["yolo_frames"] == FRAMES  # the alert clip's frames are still there


@needs_ffmpeg
def test_a_publish_over_budget_is_refused_with_the_estimate(client, s3client, labeled, monkeypatch):
    # D8: frames and bytes are estimated before anything is queued
    adm, cid = labeled["adm"], labeled["cid"]
    before = _all_keys(s3client)
    monkeypatch.setattr(tagging, "MAX_FRAMES", 50)
    r = _publish(client, adm, cid)
    assert r.status_code == 400 and "frames" in r.json()["detail"] and "about" in r.json()["detail"], r.text
    monkeypatch.setattr(tagging, "MAX_FRAMES", 300_000)
    monkeypatch.setattr(tagging, "MAX_BYTES", 1000)
    r = _publish(client, adm, cid)
    assert r.status_code == 400 and "GB" in r.json()["detail"], r.text
    assert _all_keys(s3client) == before
    with session_scope(client.app.state.engine) as s:
        assert s.scalar(select(m.TaggingPublish.id)) is None
    assert tagging.MAX_FRAMES == 300_000 and tagging.MAX_BYTES == 1000


@needs_ffmpeg
def test_publish_counts_update_while_it_runs(client, s3client, labeled, monkeypatch):
    # D8: PublishOut counts follow the job, clip by clip
    real_clip = tagging._Batch.clip
    seen = []

    def clip(self, info):
        with session_scope(client.app.state.engine) as s:
            row = s.get(m.TaggingPublish, self.pub.id)
            seen.append((row.state, row.tasks, row.yolo_frames))
        return real_clip(self, info)

    monkeypatch.setattr(tagging._Batch, "clip", clip)
    assert _publish(client, labeled["adm"], labeled["cid"]).json()["state"] == "ready"
    assert seen == [("running", 0, 0), ("running", 1, FRAMES), ("running", 2, 2 * FRAMES)]
